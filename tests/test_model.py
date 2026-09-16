import tempfile
import unittest
import pandas as pd
from openpyxl import load_workbook
from macro_forecast.synthetic import build_synthetic_dataset
from macro_forecast.validation import DataValidationError
from macro_forecast.model import (fit_model, rolling_backtest, forecast_scenarios, summarize,
                                  benchmark_backtest, error_metrics, variance_inflation_factors)
from macro_forecast.reporting import generate_demo_artifacts, REPORT_SHEETS


class MacroForecastTest(unittest.TestCase):
    def setUp(self): self.data = build_synthetic_dataset()
    def test_repeatable_and_complete(self):
        other=build_synthetic_dataset(); self.assertEqual(len(self.data.history),60); pd.testing.assert_frame_equal(self.data.history,other.history)
    def test_model_has_expected_coefficients(self): self.assertEqual(len(fit_model(self.data.history).coefficients),9)
    def test_r_squared_is_bounded(self): self.assertTrue(0 <= fit_model(self.data.history).r_squared <= 1)
    def test_fitted_reconciles(self):
        f=fit_model(self.data.history).fitted; self.assertLess((f.revenue_mn-f.fitted_revenue_mn-f.residual_mn).abs().max(),1e-8)
    def test_backtest_has_24_months(self): self.assertEqual(len(rolling_backtest(self.data.history)),24)
    def test_mape_is_reasonable(self): self.assertLess(summarize(self.data.history,self.data.scenarios)["mape"],0.08)
    def test_error_metrics_include_mae_rmse_and_mape(self):
        metrics=error_metrics(rolling_backtest(self.data.history)); self.assertEqual(set(metrics),{"mae","rmse","mape"}); self.assertGreater(metrics["rmse"],0)
    def test_benchmark_uses_same_holdout_period(self):
        self.assertEqual(len(benchmark_backtest(self.data.history)),len(rolling_backtest(self.data.history)))
    def test_model_beats_seasonal_naive_benchmark(self):
        result=summarize(self.data.history,self.data.scenarios); self.assertLess(result["metrics"]["rmse"],result["benchmark_metrics"]["rmse"])
    def test_vif_covers_all_economic_drivers(self):
        vif=variance_inflation_factors(self.data.history); self.assertEqual(len(vif),5); self.assertTrue((vif["vif"]>=1).all())
    def test_scenarios_have_36_rows(self): self.assertEqual(len(forecast_scenarios(self.data.history,self.data.scenarios)),36)
    def test_scenario_direction(self):
        t=summarize(self.data.history,self.data.scenarios)["scenario_totals"].set_index("scenario"); self.assertGreater(t.loc["乐观","forecast_revenue_mn"],t.loc["基准","forecast_revenue_mn"]); self.assertLess(t.loc["压力","forecast_revenue_mn"],t.loc["基准","forecast_revenue_mn"])
    def test_duplicate_month_rejected(self):
        x=pd.concat([self.data.history,self.data.history.iloc[[0]]]);
        with self.assertRaisesRegex(DataValidationError,"月份不能重复"): fit_model(x)
    def test_missing_field_rejected(self):
        with self.assertRaisesRegex(DataValidationError,"缺少必要字段"): fit_model(self.data.history.drop(columns="inflation"))
    def test_nonpositive_revenue_rejected(self):
        x=self.data.history.copy(); x.loc[0,"revenue_mn"]=0
        with self.assertRaisesRegex(DataValidationError,"必须为正数"): fit_model(x)
    def test_invalid_backtest_window_rejected(self):
        with self.assertRaisesRegex(DataValidationError,"窗口不合理"): rolling_backtest(self.data.history,9)
    def test_duplicate_scenario_key_rejected(self):
        x=pd.concat([self.data.scenarios,self.data.scenarios.iloc[[0]]])
        with self.assertRaisesRegex(DataValidationError,"组合不能重复"): forecast_scenarios(self.data.history,x)
    def test_workbooks_reopen(self):
        with tempfile.TemporaryDirectory() as d:
            i,r=generate_demo_artifacts(d); wi=load_workbook(i); wr=load_workbook(r)
            self.assertEqual(wi.sheetnames,["历史数据","情景假设","使用说明"]); self.assertEqual(wr.sheetnames,REPORT_SHEETS)
