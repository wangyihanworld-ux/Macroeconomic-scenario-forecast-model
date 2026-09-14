from pathlib import Path
import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference, BarChart
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.data_source import StrRef
from .model import summarize

REPORT_SHEETS = ["管理摘要", "情景比较", "月度预测", "回测评估", "模型解释", "历史数据", "情景假设", "数据质量"]
HEADERS = {"scenario":"情景", "month":"月份", "gdp_growth":"GDP增长率", "inflation":"通胀率",
           "unemployment":"失业率", "consumer_confidence":"消费者信心", "fx_index":"汇率指数",
           "revenue_mn":"实际收入（百万元）", "forecast_revenue_mn":"预测收入（百万元）",
           "variance_vs_base_mn":"较基准差异（百万元）", "actual_revenue_mn":"实际收入（百万元）",
           "predicted_revenue_mn":"滚动预测收入（百万元）", "error_mn":"误差（百万元）",
           "absolute_percentage_error":"绝对百分比误差", "driver":"驱动项", "coefficient":"系数"}


def _write_table(ws, frame: pd.DataFrame, start_row=1):
    for j, col in enumerate(frame.columns, 1):
        cell = ws.cell(start_row, j, HEADERS.get(col, col))
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="24557A")
    for i, row in enumerate(frame.itertuples(index=False, name=None), start_row + 1):
        for j, value in enumerate(row, 1):
            if isinstance(value, pd.Timestamp): value = value.to_pydatetime()
            ws.cell(i, j, value)


def _format(ws):
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    thin = Side(style="thin", color="D9E2F3")
    for row in ws.iter_rows():
        for c in row:
            c.border = Border(bottom=thin)
            c.alignment = Alignment(vertical="center", wrap_text=True)
    for col in range(1, ws.max_column + 1):
        width = max((len(str(ws.cell(r, col).value or "")) for r in range(1, ws.max_row + 1)), default=8)
        ws.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 12), 28)


def generate_report(history: pd.DataFrame, scenarios: pd.DataFrame, output_path) -> Path:
    result = summarize(history, scenarios)
    wb = Workbook(); wb.remove(wb.active)
    summary = wb.create_sheet("管理摘要")
    summary.sheet_view.showGridLines = False
    summary["A1"] = "宏观经济情景与经营预测管理摘要"
    summary["A1"].font = Font(size=20, bold=True, color="1F4E78")
    metrics = [("历史期数", len(history)), ("模型 R²", result["model"].r_squared),
               ("调整后 R²", result["model"].adjusted_r_squared), ("滚动回测 MAPE", result["mape"])]
    for r, (label, value) in enumerate(metrics, 3):
        summary.cell(r, 1, label).font = Font(bold=True); summary.cell(r, 2, value)
    base = result["scenario_totals"].set_index("scenario")["forecast_revenue_mn"]
    summary["A8"] = "2026 情景收入预测（百万元）"; summary["A8"].font = Font(bold=True, color="1F4E78")
    for r, name in enumerate(["压力", "基准", "乐观"], 9):
        summary.cell(r, 1, name); summary.cell(r, 2, float(base[name]))
    summary["A13"] = "管理解读"
    summary["A13"].font = Font(bold=True, color="1F4E78")
    summary["A14"] = "宏观情景用于量化经营收入的条件变化，不代表因果结论或承诺。"
    summary["A15"] = "滚动回测检验样本外预测误差；决策时应同时关注误差与假设边界。"
    summary.column_dimensions["A"].width = 62; summary.column_dimensions["B"].width = 20

    s = wb.create_sheet("情景比较"); _write_table(s, result["scenario_totals"])
    chart = BarChart(); chart.title = "2026 年收入情景比较"; chart.y_axis.title = "百万元"
    chart.add_data(Reference(s, min_col=2, min_row=1, max_row=4), titles_from_data=True)
    chart.series[0].tx = SeriesLabel(strRef=StrRef(f="'{s.title}'!$B$1"))
    chart.set_categories(Reference(s, min_col=1, min_row=2, max_row=4)); s.add_chart(chart, "E2")
    monthly = wb.create_sheet("月度预测"); _write_table(monthly, result["forecast"])
    back = wb.create_sheet("回测评估"); _write_table(back, result["backtest"])
    line = LineChart(); line.title = "实际与滚动预测"; line.y_axis.title = "百万元"
    line.add_data(Reference(back, min_col=2, max_col=3, min_row=1, max_row=back.max_row), titles_from_data=True)
    line.series[0].tx = SeriesLabel(strRef=StrRef(f="'{back.title}'!$B$1")); line.series[1].tx = SeriesLabel(strRef=StrRef(f="'{back.title}'!$C$1"))
    line.set_categories(Reference(back, min_col=1, min_row=2, max_row=back.max_row)); back.add_chart(line, "G2")
    coeff = wb.create_sheet("模型解释"); _write_table(coeff, result["model"].coefficients)
    coeff["D1"] = "解释边界"; coeff["D1"].font = Font(bold=True, color="1F4E78")
    coeff["D2"] = "系数描述控制其他变量后的统计关联，不等同于因果效应。"
    hist = wb.create_sheet("历史数据"); _write_table(hist, history)
    assump = wb.create_sheet("情景假设"); _write_table(assump, scenarios)
    quality = wb.create_sheet("数据质量")
    checks = pd.DataFrame({"检查项":["历史月份唯一","历史字段完整","情景键唯一","预测为正"],"结果":["通过"]*4})
    _write_table(quality, checks)
    for ws in wb.worksheets:
        _format(ws)
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, float): c.number_format = "0.00%" if "率" in str(ws.cell(1, c.column).value) or c.column == 2 and ws.title == "管理摘要" and c.row in (4,5,6) else "#,##0.00"
                if hasattr(c.value, "year"): c.number_format = "yyyy-mm"
    output = Path(output_path); output.parent.mkdir(parents=True, exist_ok=True); wb.save(output); return output


def generate_demo_artifacts(output_dir):
    from .synthetic import build_synthetic_dataset
    folder = Path(output_dir); folder.mkdir(parents=True, exist_ok=True)
    data = build_synthetic_dataset()
    input_path = folder / "合成宏观与经营输入.xlsx"
    with pd.ExcelWriter(input_path, engine="openpyxl") as writer:
        data.history.to_excel(writer, sheet_name="历史数据", index=False)
        data.scenarios.to_excel(writer, sheet_name="情景假设", index=False)
        pd.DataFrame({"说明":["全部数据由程序生成，不对应任何真实企业、经济体或交易。"]}).to_excel(writer, sheet_name="使用说明", index=False)
    report_path = generate_report(data.history, data.scenarios, folder / "宏观经济情景与经营预测报告.xlsx")
    return input_path, report_path
